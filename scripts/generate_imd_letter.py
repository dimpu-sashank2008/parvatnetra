# -*- coding: utf-8 -*-
"""
scripts/generate_imd_letter.py
==============================
Generates the official institutional IMD API Recommendation & Permission Letter
as both a Microsoft Word (.docx) file and an A4-printable HTML document.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml


def generate_docx(output_path: str):
    doc = docx.Document()

    # Set standard margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(10.5)
    style.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # Header Notice
    p_top = doc.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_top = p_top.add_run('[ON THE OFFICIAL LETTERHEAD OF THE INSTITUTE]')
    run_top.bold = True
    run_top.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    run_top.font.size = Pt(10)

    # Ref and Date line
    p_ref = doc.add_paragraph()
    p_ref.paragraph_format.space_before = Pt(10)
    p_ref.paragraph_format.space_after = Pt(14)
    run_ref = p_ref.add_run('Ref. No.: ___________________________')
    run_ref.font.size = Pt(10)
    run_date = p_ref.add_run('\t\t\t\tDate: ___________________')
    run_date.font.size = Pt(10)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(14)
    run_title = p_title.add_run('PERMISSION / RECOMMENDATION FOR ACCESS TO IMD APIs')
    run_title.bold = True
    run_title.font.size = Pt(12.5)

    # Addressee
    p_to = doc.add_paragraph()
    p_to.paragraph_format.space_after = Pt(10)
    r = p_to.add_run('To\n')
    r.bold = True
    p_to.add_run('Director General of Meteorology\n')
    p_to.add_run('India Meteorological Department (IMD)\n')
    p_to.add_run('Ministry of Earth Sciences, Government of India\n')
    p_to.add_run('Lodhi Road, New Delhi – 110003')

    # Certification Body
    p_cert = doc.add_paragraph()
    p_cert.paragraph_format.space_after = Pt(8)
    p_cert.paragraph_format.line_spacing = 1.15
    p_cert.add_run('This is to certify that ')
    r_names = p_cert.add_run('[Name(s) of Student(s) / Research Scholar(s)]')
    r_names.bold = True
    p_cert.add_run(', of ')
    p_cert.add_run('[Department / Programme: e.g. Dept. of Computer Science & Engineering / Civil Engineering]')
    p_cert.add_run(', ')
    r_inst = p_cert.add_run('[Institute / University Name]')
    r_inst.bold = True
    p_cert.add_run(', is/are undertaking the project titled ')
    r_proj = p_cert.add_run('“PARVAT NETRA: Predictive AI & Hillslope Telemetry Early-Warning System for the North-Eastern Region (PAHAD AI)”')
    r_proj.bold = True
    p_cert.add_run(' as a part of ')
    p_cert.add_run('[Smart India Hackathon / Academic Research / Capstone Project]')
    p_cert.add_run(' under the supervision of ')
    r_guide = p_cert.add_run('[Name & Designation of Faculty / Project Guide]')
    r_guide.bold = True
    p_cert.add_run('.')

    p_lead = doc.add_paragraph('The project details and requirement for IMD API access are as follows:')
    p_lead.paragraph_format.space_after = Pt(6)

    # Table
    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(1.8), Inches(4.7)]
    for row in table.rows:
        for i, w in enumerate(col_widths):
            row.cells[i].width = w

    table_data = [
        ('Project Objective',
         'To develop an autonomous, physics-grounded hillslope disaster intelligence and early-warning platform (PARVAT NETRA / PAHAD AI) for the North-Eastern Region of India (NER). The system correlates real-time precipitation, geotechnical Factor of Safety (FoS), and InSAR deformation to deliver early warnings with up to 24-hour lead time across critical national highway corridors.'),
        ('IMD APIs Required',
         '1. Real-time Automated Weather Station (AWS) 15-minute precipitation & temperature telemetry.\n'
         '2. High-resolution Gridded Daily & Hourly Rainfall Products (0.25° x 0.25° resolution).\n'
         '3. Doppler Weather Radar (DWR) composite reflectivity & nowcasting feeds for NER (Agartala, Mohanbari, Sohra radars).\n'
         '4. District-level Heavy Rainfall & Flash Flood Advisories.'),
        ('Inputs',
         'IMD real-time & gridded rainfall telemetry + GSI 1:250,000 geological quadrangle baselines, Sentinel-1 InSAR ground deformation velocities, CWC hydrometric river levels, and in-situ IoT piezometer/inclinometer data.'),
        ('Methodology',
         'Mohr-Coulomb Infinite Slope stability physics integrated with calibrated Gradient Boosted event classifiers, Antecedent Precipitation Index (API-24h, API-72h, API-30d) threshold curves, transient groundwater seepage modeling, and 2-of-3 multimodal confirmation gating before alert recommendation.'),
        ('Expected Outputs',
         'Real-time Composite Risk Index (CRI) geospatial command dashboard, early-warning alerts with quantified lead time, OASIS CAP v1.2 emergency feeds, automated tactical evacuation routing for SDRF/NDRF, and multilingual emergency voice broadcasts.'),
        ('Project Timeline',
         'October 2024 to September 2026 (Standing National Disaster-Intelligence Research & Implementation)')
    ]

    for row_idx, (particular, details) in enumerate(table_data):
        cell_p = table.cell(row_idx, 0)
        cell_d = table.cell(row_idx, 1)

        cp_p = cell_p.paragraphs[0]
        cp_p.paragraph_format.space_after = Pt(2)
        r_p = cp_p.add_run(particular)
        r_p.bold = True
        r_p.font.size = Pt(9.5)

        cp_d = cell_d.paragraphs[0]
        cp_d.paragraph_format.space_after = Pt(2)
        r_d = cp_d.add_run(details)
        r_d.font.size = Pt(9.5)

    # Style borders for table
    tblPr = table._tbl.tblPr
    borders = parse_xml(r'''
        <w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
            <w:top w:val="single" w:sz="4" w:space="0" w:color="9CA3AF"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="9CA3AF"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="9CA3AF"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="9CA3AF"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="D1D5DB"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="D1D5DB"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

    # Undertakings
    p_rec = doc.add_paragraph()
    p_rec.paragraph_format.space_before = Pt(10)
    p_rec.paragraph_format.space_after = Pt(6)
    p_rec.paragraph_format.line_spacing = 1.15
    p_rec.add_run('The Institute hereby grants permission and recommends the above-mentioned student(s)/team for obtaining access to the required IMD APIs, solely for the purpose of the stated project and for the duration indicated above.')

    p_und = doc.add_paragraph()
    p_und.paragraph_format.space_after = Pt(12)
    p_und.paragraph_format.line_spacing = 1.15
    p_und.add_run('The Institute undertakes that the IMD APIs, if access is granted, shall be used only for the stated project; API credentials/data shall not be shared with any unauthorized person or third party; and all terms, conditions, usage limits and data-security requirements prescribed by IMD shall be strictly complied with. The data shall not be redistributed or used for commercial purposes in any case.')

    # Signatory block
    p_sig_title = doc.add_paragraph()
    r_sig = p_sig_title.add_run('Authorized Signatory')
    r_sig.bold = True
    p_sig_title.paragraph_format.space_after = Pt(4)

    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.line_spacing = 1.2
    p_sig.paragraph_format.space_after = Pt(20)
    p_sig.add_run('Name: _________________________________________________\n')
    p_sig.add_run('Designation: Head of Institute / Director / Dean / Principal\n')
    p_sig.add_run('Institute/University: _____________________________________\n')
    p_sig.add_run('Official Email: _________________________________________\n')
    p_sig.add_run('Contact No.: ___________________________________________')

    p_seal = doc.add_paragraph()
    r_seal = p_seal.add_run('Signature & Official Seal')
    r_seal.bold = True

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    print(f"Generated Word Document: {output_path}")


def generate_html(output_path: str):
    html = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>IMD API Access Recommendation Letter - PARVAT NETRA</title>
<style>
  @page {
    size: A4;
    margin: 20mm 15mm 20mm 15mm;
  }
  body {
    font-family: "Calibri", "Segoe UI", Arial, sans-serif;
    color: #111827;
    line-height: 1.35;
    margin: 0;
    padding: 0;
    background: #f9fafb;
  }
  .page-container {
    max-width: 210mm;
    margin: 20px auto;
    background: #ffffff;
    padding: 25mm 20mm 20mm 20mm;
    box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.1);
    box-sizing: border-box;
  }
  @media print {
    body { background: transparent; }
    .page-container {
      margin: 0;
      padding: 0;
      box-shadow: none;
      width: 100%;
    }
    .no-print { display: none; }
  }
  .header-note {
    text-align: center;
    font-size: 11pt;
    font-weight: bold;
    color: #6b7280;
    margin-bottom: 25px;
    letter-spacing: 0.5px;
  }
  .ref-date-row {
    display: flex;
    justify-content: space-between;
    font-size: 10.5pt;
    margin-bottom: 25px;
  }
  .doc-title {
    text-align: center;
    font-size: 13pt;
    font-weight: bold;
    text-decoration: underline;
    margin-bottom: 20px;
    color: #111827;
  }
  .addressee {
    font-size: 10.5pt;
    margin-bottom: 18px;
    line-height: 1.4;
  }
  .certification-body {
    font-size: 10.5pt;
    text-align: justify;
    margin-bottom: 14px;
    line-height: 1.45;
  }
  table.particulars-table {
    width: 100%;
    border-collapse: collapse;
    margin: 14px 0 16px 0;
    font-size: 9.5pt;
  }
  table.particulars-table th, table.particulars-table td {
    border: 1px solid #9ca3af;
    padding: 7px 10px;
    vertical-align: top;
  }
  table.particulars-table th {
    background-color: #f3f4f6;
    text-align: left;
    font-weight: bold;
    color: #111827;
  }
  .undertaking {
    font-size: 10pt;
    text-align: justify;
    margin-bottom: 14px;
    line-height: 1.4;
  }
  .signatory-block {
    margin-top: 25px;
    font-size: 10.5pt;
    line-height: 1.5;
  }
  .signature-space {
    margin-top: 40px;
    font-weight: bold;
  }
  .btn-print {
    background: #2563eb;
    color: #ffffff;
    border: none;
    padding: 10px 20px;
    font-size: 14px;
    font-weight: bold;
    border-radius: 6px;
    cursor: pointer;
    box-shadow: 0 2px 4px rgba(0,0,0,0.1);
    margin-bottom: 20px;
  }
</style>
</head>
<body>

<div style="text-align: center; padding-top: 20px;" class="no-print">
  <button class="btn-print" onclick="window.print()">🖨️ Print onto Official Letterhead / Save as PDF</button>
</div>

<div class="page-container">
  <div class="header-note">[ON THE OFFICIAL LETTERHEAD OF THE INSTITUTE]</div>

  <div class="ref-date-row">
    <div>Ref. No.: ___________________________</div>
    <div>Date: ___________________</div>
  </div>

  <div class="doc-title">PERMISSION / RECOMMENDATION FOR ACCESS TO IMD APIs</div>

  <div class="addressee">
    <strong>To</strong><br>
    Director General of Meteorology<br>
    India Meteorological Department (IMD)<br>
    Ministry of Earth Sciences, Government of India<br>
    Lodhi Road, New Delhi – 110003
  </div>

  <div class="certification-body">
    This is to certify that <strong>[Name(s) of Student(s)/Research Scholar(s)]</strong>, of <strong>[Department/Programme]</strong>, <strong>[Institute/University Name]</strong>, is/are undertaking the project titled <strong>“PARVAT NETRA: Predictive AI & Hillslope Telemetry Early-Warning System for the North-Eastern Region (PAHAD AI)”</strong> as a part of <strong>[Smart India Hackathon / Academic Project / Research Project / Dissertation etc.]</strong> under the supervision of <strong>[Name & Designation of Faculty/Project Guide]</strong>.
  </div>

  <div style="font-size: 10pt; font-weight: bold; margin-bottom: 6px;">
    The project details and requirement for IMD API access are as follows:
  </div>

  <table class="particulars-table">
    <thead>
      <tr>
        <th style="width: 25%;">Particular</th>
        <th style="width: 75%;">Details</th>
      </tr>
    </thead>
    <tbody>
      <tr>
        <td><strong>Project Objective</strong></td>
        <td>To develop an autonomous, physics-grounded hillslope disaster intelligence and early-warning platform (PARVAT NETRA / PAHAD AI) for the North-Eastern Region of India (NER). The system correlates real-time precipitation, geotechnical Factor of Safety (FoS), and InSAR deformation to deliver early warnings with up to 24-hour lead time across critical national highway corridors.</td>
      </tr>
      <tr>
        <td><strong>IMD APIs Required</strong></td>
        <td>
          1. Real-time Automated Weather Station (AWS) 15-minute precipitation & temperature telemetry.<br>
          2. High-resolution Gridded Daily and Hourly Rainfall Products (0.25° x 0.25° resolution).<br>
          3. Doppler Weather Radar (DWR) composite reflectivity and nowcasting feeds for NER (Agartala, Mohanbari, Sohra radars).<br>
          4. District-level Heavy Rainfall & Flash Flood Advisories.
        </td>
      </tr>
      <tr>
        <td><strong>Inputs</strong></td>
        <td>IMD real-time & gridded rainfall telemetry + GSI 1:250,000 geological quadrangle baselines, Sentinel-1 InSAR ground deformation velocities, CWC hydrometric river levels, and in-situ IoT piezometer/inclinometer data.</td>
      </tr>
      <tr>
        <td><strong>Methodology</strong></td>
        <td>Mohr-Coulomb Infinite Slope stability physics integrated with calibrated Gradient Boosted event classifiers, Antecedent Precipitation Index (API-24h, API-72h, API-30d) threshold curves, transient groundwater seepage modeling, and 2-of-3 multimodal confirmation gating before alert recommendation.</td>
      </tr>
      <tr>
        <td><strong>Expected Outputs</strong></td>
        <td>Real-time Composite Risk Index (CRI) geospatial command dashboard, early-warning alerts with quantified lead time, OASIS CAP v1.2 emergency feeds, automated tactical evacuation routing for SDRF/NDRF, and multilingual emergency voice broadcasts.</td>
      </tr>
      <tr>
        <td><strong>Project Timeline</strong></td>
        <td>October 2024 to September 2026 (Standing National Disaster-Intelligence Research & Implementation)</td>
      </tr>
    </tbody>
  </table>

  <div class="undertaking">
    The Institute hereby grants permission and recommends the above-mentioned student(s)/team for obtaining access to the required IMD APIs, solely for the purpose of the stated project and for the duration indicated above.
  </div>

  <div class="undertaking">
    The Institute undertakes that the IMD APIs, if access is granted, shall be used only for the stated project; API credentials/data shall not be shared with any unauthorized person or third party; and all terms, conditions, usage limits and data-security requirements prescribed by IMD shall be strictly complied with. The data shall not be redistributed or used for commercial purposes in any case.
  </div>

  <div class="signatory-block">
    <strong>Authorized Signatory</strong><br><br>
    Name: _________________________________________________<br>
    Designation: Head of Institute / Director / Dean / Principal<br>
    Institute/University: _____________________________________<br>
    Official Email: _________________________________________<br>
    Contact No.: ___________________________________________<br>
    <div class="signature-space">Signature & Official Seal</div>
  </div>
</div>

</body>
</html>"""
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(html)
    print(f"Generated HTML Document: {output_path}")


if __name__ == "__main__":
    docx_out = r"c:\Users\dimpu\Downloads\PARVAT_NETRA_PAHAD_AI_FIRST\silly-fermi\docs\PARVAT_NETRA_IMD_PERMISSION_LETTER.docx"
    html_out = r"c:\Users\dimpu\Downloads\PARVAT_NETRA_PAHAD_AI_FIRST\silly-fermi\docs\PARVAT_NETRA_IMD_PERMISSION_LETTER.html"
    generate_docx(docx_out)
    generate_html(html_out)
